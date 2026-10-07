"""Export clean_bengkel ke file Excel siap kirim ke klien."""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from db import connect

COLUMNS = [
    ("Nama Bengkel", "nama"), ("Alamat", "alamat"), ("No HP", "no_hp"),
    ("Kecamatan", "kecamatan"), ("Kabupaten/Kota", "kabupaten"), ("Provinsi", "provinsi"),
    ("Latitude", "latitude"), ("Longitude", "longitude"), ("Kategori", "kategori"),
]
# Koordinat hanya ditulis kalau lolos validasi; titik yang jelas salah dikosongkan.


def main() -> None:
    parser = argparse.ArgumentParser()
    default = Path(__file__).resolve().parent.parent / "export" / f"bengkel_{date.today():%Y%m%d}.xlsx"
    parser.add_argument("--out", type=Path, default=default)
    args = parser.parse_args()

    with connect() as conn, conn.cursor() as cur:
        cur.execute(
            f"""select {", ".join(src for _, src in COLUMNS)}, koordinat_valid
                from clean_bengkel
                order by provinsi, kabupaten, kecamatan, nama"""
        )
        rows = []
        lat_i = next(i for i, (_, src) in enumerate(COLUMNS) if src == "latitude")
        lon_i = next(i for i, (_, src) in enumerate(COLUMNS) if src == "longitude")
        for record in cur.fetchall():
            values = list(record[:-1])
            if not record[-1]:  # koordinat_valid False
                values[lat_i] = None
                values[lon_i] = None
            rows.append(values)

    wb = Workbook()
    ws = wb.active
    ws.title = "Bengkel"
    header_fill = PatternFill("solid", fgColor="1F3A4D")
    header_font = Font(bold=True, color="FFFFFF")
    for col, (label, _) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    phone_col = next(i for i, (_, src) in enumerate(COLUMNS, start=1) if src == "no_hp")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == phone_col and value is not None:
                cell.number_format = "@"  # teks, biar angka 0 di depan tidak hilang

    widths = [30, 40, 16, 20, 24, 16, 12, 12, 10]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNS)).column_letter}{len(rows) + 1}"

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    dengan_hp = sum(1 for row in rows if row[2])
    print(f"export: {len(rows)} baris ke {args.out}")
    if rows:
        print(f"  punya no HP: {dengan_hp} ({dengan_hp * 100 // len(rows)}%)")


if __name__ == "__main__":
    main()
